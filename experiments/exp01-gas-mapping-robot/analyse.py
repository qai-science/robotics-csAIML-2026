"""Analyse a gas-mapping run: reads the base station's Excel log and the robot's truth file.

Writes to results/:
    fig1_true_gas_field.png     true CH4 and CO2 at sensor height, and both blended in one colour map
    fig2_vertical_profiles.png  side view through each leak: CH4 rises, CO2 pools at the floor
    fig3_cloud_3d.png           3-D gas clouds (point density = concentration) with the robot's path
    fig4_readings_on_path.png   what the robot measured, where (colour = ppm, black rings = alarms)
    fig5_measured_vs_true.png   map rebuilt from the readings vs the true map
    fig6_time_series.png        measured vs true concentration over time (shows sensor lag)
    summary.json                all numbers quoted in the README

    python analyse.py
"""
import csv
import json
import math
import os

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
from openpyxl import load_workbook

import gas_field as gf

HERE = os.path.dirname(os.path.abspath(__file__))
RES = os.path.join(HERE, 'results')
GASES = ('CH4', 'CO2')
NAMES = {'CH4': 'Methane (CH$_4$)', 'CO2': 'Carbon dioxide (CO$_2$)'}


def cmap(gas):
    return LinearSegmentedColormap.from_list(gas, [(0.96, 0.96, 0.96), gf.COLOR[gas], tuple(0.45 * c for c in gf.COLOR[gas])])


def load_excel(path):
    ws = load_workbook(path, read_only=True)['Readings']
    rows = list(ws.iter_rows(values_only=True))
    head = rows[0]
    return [dict(zip(head, r)) for r in rows[1:]]


def load_truth(path):
    with open(path) as f:
        return [{k: float(v) for k, v in r.items()} for r in csv.DictReader(f)]


def true_grid(gas, z, n=101):
    xs = np.linspace(-2.5, 2.5, n)
    X, Y = np.meshgrid(xs, xs)
    C = np.vectorize(lambda x, y: gf.concentration(gas, x, y, z))(X, Y)
    return X, Y, C


def blend(c1, c2):
    """RGB image mixing the two gas colours by their normalised concentrations."""
    a = np.clip(c1 / gf.ALARM['CH4'], 0, 1)[..., None]
    b = np.clip(c2 / gf.ALARM['CO2'], 0, 1)[..., None]
    white = np.ones(3)
    g1, g2 = np.array(gf.COLOR['CH4']), np.array(gf.COLOR['CO2'])
    return np.clip(white * (1 - a) * (1 - b) + g1 * a * (1 - b) + g2 * b * (1 - a) + (g1 + g2) / 2 * a * b, 0, 1)


def idw(px, py, pv, X, Y, power=2.0, radius=0.6):
    """Inverse-distance-weighted map from scattered readings; cells far from any reading are NaN."""
    out = np.full(X.shape, np.nan)
    pts = np.c_[px, py]
    for i in range(X.shape[0]):
        for j in range(X.shape[1]):
            d = np.hypot(pts[:, 0] - X[i, j], pts[:, 1] - Y[i, j])
            m = d < radius
            if m.any():
                w = 1 / np.maximum(d[m], 0.02) ** power
                out[i, j] = (w * pv[m]).sum() / w.sum()
    return out


def lag_seconds(t, true, meas, max_lag=6.0):
    """Delay that best lines up the measured series with the true one (cross-correlation)."""
    dt = t[1] - t[0]
    a, b = true - true.mean(), meas - meas.mean()
    best, best_k = -np.inf, 0
    for k in range(int(max_lag / dt) + 1):
        c = np.dot(a[:len(a) - k], b[k:]) if k else np.dot(a, b)
        if c > best:
            best, best_k = c, k
    return best_k * dt


def main():
    xl = load_excel(os.path.join(RES, 'gas_log.xlsx'))
    tr = load_truth(os.path.join(RES, 'robot_truth.csv'))
    t = np.array([r['t'] for r in tr])
    px, py = np.array([r['x'] for r in tr]), np.array([r['y'] for r in tr])
    z = float(np.median([r['z'] for r in tr]))
    summary = {'readings_in_excel': len(xl), 'readings_sent': len(tr), 'duration_s': float(t[-1]),
               'sensor_height_m': round(z, 3),
               'mean_link_delay_ms': round(float(np.mean([r['Link delay (ms)'] for r in xl])), 2),
               'max_link_delay_ms': round(float(np.max([r['Link delay (ms)'] for r in xl])), 2),
               'path_length_m': round(float(np.sum(np.hypot(np.diff(px), np.diff(py)))), 2)}

    # ---- Fig 1: true fields at sensor height + blended map
    grids = {g: true_grid(g, z) for g in GASES}
    fig, axes = plt.subplots(1, 3, figsize=(16, 5))
    for ax, g in zip(axes, GASES):
        X, Y, C = grids[g]
        im = ax.pcolormesh(X, Y, C, cmap=cmap(g), shading='auto')
        ax.contour(X, Y, C, levels=[gf.WARN[g], gf.ALARM[g]], colors=['#b8860b', '#c00000'], linewidths=1.2)
        s = gf.SOURCES[g]
        ax.plot(s['x'], s['y'], 'k*', ms=12)
        fig.colorbar(im, ax=ax, label='ppm')
        ax.set_title(f'{NAMES[g]} at {z:.2f} m height\n(yellow = warning, red = alarm contour, ★ = leak)', fontsize=10)
        ax.set_aspect('equal')
    axes[2].imshow(blend(grids['CH4'][2] - 2, grids['CO2'][2] - 420), origin='lower', extent=[-2.5, 2.5, -2.5, 2.5])
    axes[2].plot(px, py, color='k', lw=0.6, alpha=0.6)
    axes[2].set_title('Both gases blended (green = CH$_4$, purple = CO$_2$)\nwith the robot path', fontsize=10)
    for ax in axes:
        ax.set_xlabel('x (m)')
        ax.set_ylabel('y (m)')
    fig.tight_layout()
    fig.savefig(os.path.join(RES, 'fig1_true_gas_field.png'), dpi=140)
    plt.close(fig)

    # ---- Fig 2: vertical sections through each leak
    fig, axes = plt.subplots(1, 2, figsize=(13, 4))
    for ax, g in zip(axes, GASES):
        s = gf.SOURCES[g]
        xs, zs = np.linspace(-2.5, 2.5, 120), np.linspace(0, 1.6, 60)
        XX, ZZ = np.meshgrid(xs, zs)
        C = np.vectorize(lambda x, zz: gf.concentration(g, x, s['y'], zz))(XX, ZZ)
        im = ax.pcolormesh(XX, ZZ, C, cmap=cmap(g), shading='auto')
        ax.axhline(z, color='k', ls='--', lw=1)
        ax.text(-2.4, z + 0.04, 'sensor height', fontsize=8)
        fig.colorbar(im, ax=ax, label='ppm')
        ax.set_title(f'{NAMES[g]}: side view through the leak (y = {s["y"]} m)', fontsize=10)
        ax.set_xlabel('x (m)  →  draught direction')
        ax.set_ylabel('height z (m)')
    fig.tight_layout()
    fig.savefig(os.path.join(RES, 'fig2_vertical_profiles.png'), dpi=140)
    plt.close(fig)

    # ---- Fig 3: 3-D clouds
    fig = plt.figure(figsize=(9, 7))
    ax = fig.add_subplot(111, projection='3d')
    for g, seed in (('CH4', 11), ('CO2', 22)):
        pts = np.array(gf.sample_cloud(g, 2500, seed=seed))
        ax.scatter(pts[:, 0], pts[:, 1], pts[:, 2], s=3, alpha=0.25, color=gf.COLOR[g], label=NAMES[g])
    ax.plot(px, py, np.full_like(px, z), color='k', lw=1, label='robot sensor path')
    ax.set_xlim(-2.5, 2.5)
    ax.set_ylim(-2.5, 2.5)
    ax.set_zlim(0, 1.6)
    ax.set_xlabel('x (m)')
    ax.set_ylabel('y (m)')
    ax.set_zlabel('z (m)')
    ax.view_init(elev=22, azim=-60)
    ax.legend(loc='upper left', fontsize=8)
    ax.set_title('Gas clouds as drawn in CoppeliaSim (density = concentration)', fontsize=10)
    fig.tight_layout()
    fig.savefig(os.path.join(RES, 'fig3_cloud_3d.png'), dpi=140)
    plt.close(fig)

    # ---- Fig 4: readings on the path (from the Excel file)
    fig, axes = plt.subplots(1, 2, figsize=(12, 5))
    ex = np.array([r['X (m)'] for r in xl])
    ey = np.array([r['Y (m)'] for r in xl])
    for ax, g in zip(axes, GASES):
        v = np.array([r[f'{g} (ppm)'] for r in xl], dtype=float)
        sc = ax.scatter(ex, ey, c=v, cmap=cmap(g), s=18, edgecolors='none')
        al = np.array([r[f'{g} status'] == 'ALARM' for r in xl])
        ax.scatter(ex[al], ey[al], s=40, facecolors='none', edgecolors='k', linewidths=0.8, label='ALARM reading')
        s = gf.SOURCES[g]
        ax.plot(s['x'], s['y'], 'r*', ms=12, label='true leak')
        fig.colorbar(sc, ax=ax, label='measured ppm')
        ax.set_xlim(-2.5, 2.5)
        ax.set_ylim(-2.5, 2.5)
        ax.set_aspect('equal')
        ax.legend(fontsize=8, loc='lower right')
        ax.set_title(f'{NAMES[g]}: readings logged at the base station', fontsize=10)
    fig.tight_layout()
    fig.savefig(os.path.join(RES, 'fig4_readings_on_path.png'), dpi=140)
    plt.close(fig)

    # ---- Fig 5: map rebuilt from readings vs truth, and accuracy numbers
    xs = np.linspace(-2.5, 2.5, 51)
    X, Y = np.meshgrid(xs, xs)
    fig, axes = plt.subplots(2, 3, figsize=(15, 9))
    for row, g in enumerate(GASES):
        meas = np.array([r[f'{g}_meas'] for r in tr])
        true = np.array([r[f'{g}_true'] for r in tr])
        M = idw(px, py, meas, X, Y)
        T = np.vectorize(lambda x, y: gf.concentration(g, x, y, z))(X, Y)
        vmax = max(np.nanmax(M), T.max())
        for ax, data, title in ((axes[row, 0], T, 'true'), (axes[row, 1], M, 'rebuilt from readings')):
            im = ax.pcolormesh(X, Y, data, cmap=cmap(g), vmin=gf.SOURCES[g]['background'], vmax=vmax, shading='auto')
            ax.set_title(f'{NAMES[g]}: {title}', fontsize=10)
            ax.set_aspect('equal')
            fig.colorbar(im, ax=ax, label='ppm')
        im = axes[row, 2].pcolormesh(X, Y, M - T, cmap='RdBu_r', vmin=-1500, vmax=1500, shading='auto')
        axes[row, 2].plot(px, py, 'k', lw=0.4, alpha=0.5)
        axes[row, 2].set_title(f'{NAMES[g]}: rebuilt - true (blank = not visited)', fontsize=10)
        axes[row, 2].set_aspect('equal')
        fig.colorbar(im, ax=axes[row, 2], label='ppm')

        k = int(np.argmax(meas))
        i, j = np.unravel_index(np.nanargmax(M), M.shape)
        ti, tj = np.unravel_index(np.argmax(T), T.shape)
        lag = lag_seconds(t, true, meas)
        summary[g] = {
            'max_measured_ppm': float(meas.max()),
            'max_true_on_path_ppm': round(float(true.max()), 1),
            'max_reading_at': [round(float(px[k]), 2), round(float(py[k]), 2)],
            'true_hotspot_at_sensor_height': [round(float(X[ti, tj]), 2), round(float(Y[ti, tj]), 2)],
            'rebuilt_hotspot': [round(float(X[i, j]), 2), round(float(Y[i, j]), 2)],
            'hotspot_error_m': round(float(math.hypot(X[i, j] - X[ti, tj], Y[i, j] - Y[ti, tj])), 2),
            'leak_position': [gf.SOURCES[g]['x'], gf.SOURCES[g]['y']],
            'rmse_reading_vs_true_ppm': round(float(np.sqrt(np.mean((meas - true) ** 2))), 1),
            'mean_abs_error_ppm': round(float(np.mean(np.abs(meas - true))), 1),
            'sensor_time_constant_s': gf.GasSensor.SPECS[g]['tau'],
            'measured_lag_s': round(lag, 2),
            'alarm_readings': int(sum(r[f'{g} status'] == 'ALARM' for r in xl)),
            'warning_readings': int(sum(r[f'{g} status'] == 'WARNING' for r in xl)),
            'true_alarm_samples': int(np.sum(true >= gf.ALARM[g])),
            'map_coverage_pct': round(100 * float(np.mean(~np.isnan(M))), 1),
        }
    for ax in axes.flat:
        ax.set_xlabel('x (m)')
        ax.set_ylabel('y (m)')
    fig.tight_layout()
    fig.savefig(os.path.join(RES, 'fig5_measured_vs_true.png'), dpi=130)
    plt.close(fig)

    # ---- Fig 6: time series
    fig, axes = plt.subplots(2, 1, figsize=(12, 6), sharex=True)
    for ax, g in zip(axes, GASES):
        ax.plot(t, [r[f'{g}_true'] for r in tr], color='k', lw=1, label='true at the inlet')
        ax.plot(t, [r[f'{g}_meas'] for r in tr], color=gf.COLOR[g], lw=1.3, label='sensor reading')
        ax.axhline(gf.ALARM[g], color='#c00000', ls='--', lw=1, label='alarm threshold')
        ax.axhline(gf.WARN[g], color='#b8860b', ls=':', lw=1, label='warning threshold')
        ax.set_ylabel(f'{g} (ppm)')
        ax.set_title(f'{NAMES[g]}: sensor time constant {gf.GasSensor.SPECS[g]["tau"]} s, '
                     f'best-fit lag {summary[g]["measured_lag_s"]} s', fontsize=10)
        ax.legend(fontsize=8, loc='upper right')
        ax.grid(alpha=0.3)
    axes[1].set_xlabel('time (s)')
    fig.tight_layout()
    fig.savefig(os.path.join(RES, 'fig6_time_series.png'), dpi=140)
    plt.close(fig)

    with open(os.path.join(RES, 'summary.json'), 'w') as f:
        json.dump(summary, f, indent=2)
    print(json.dumps(summary, indent=2))


if __name__ == '__main__':
    main()
